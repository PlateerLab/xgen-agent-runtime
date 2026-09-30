"""Doc* (edit2docs) built-in tool family.

The engine is an optional import — every test that needs it skips cleanly
when it is not importable, so the suite stays green on minimal installs.
Deterministic paths run against the real engine with local fixtures (no
network, no LLM key).
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))

from xgen_agent_runtime.tools.base import ToolContext
from xgen_agent_runtime.tools.built_in import BUILT_IN_TOOL_CLASSES, BUILT_IN_TOOL_FEATURES
from xgen_agent_runtime.tools.built_in.doc_tools import (
    DOC_TOOL_CLASSES,
    DocAnalyzeTool,
    DocApplyEditsTool,
    DocArrangeTool,
    DocBuildTool,
    DocEditTool,
    DocGenerateTool,
    DocGuideTool,
    DocRenderTool,
    DocXmlEditTool,
    DocXmlReadTool,
)

edit2docs = pytest.importorskip("edit2docs", reason="edit2docs extra not installed")


# ── Registration ───────────────────────────────────────────


class TestRegistration:
    def test_family_registered(self):
        for name in DOC_TOOL_CLASSES:
            assert name in BUILT_IN_TOOL_CLASSES

    def test_feature_group(self):
        assert BUILT_IN_TOOL_FEATURES["documents"] == list(DOC_TOOL_CLASSES.keys())

    def test_no_browser_family(self):
        # an-web 은 제거됐다 — 서버 브라우저 도구가 다시 생기지 않는다.
        assert "browser" not in BUILT_IN_TOOL_FEATURES
        assert not any(n.startswith("Browser") for n in BUILT_IN_TOOL_CLASSES)

    def test_schemas_are_valid_shapes(self):
        for name, cls in DOC_TOOL_CLASSES.items():
            tool = cls()
            fmt = tool.to_api_format()
            assert fmt["name"] == name
            assert fmt["description"]
            assert fmt["input_schema"]["type"] == "object"


# ── Doc tools against real generated fixtures (no LLM) ─────


@pytest.fixture
def docx_path(tmp_path):
    from edit2docs.documents.docx_engine import docx_from_markdown

    data = docx_from_markdown("# Title\n\nFirst paragraph.\n\nSecond paragraph.")
    p = tmp_path / "doc.docx"
    p.write_bytes(data)
    return p


@pytest.fixture
def xlsx_path(tmp_path):
    from edit2docs.documents.xlsx_engine import xlsx_from_spec

    data = xlsx_from_spec(
        {"sheets": [{"name": "Data", "headers": ["a", "b"], "rows": [[1, 2], [3, 4]]}]}
    )
    p = tmp_path / "book.xlsx"
    p.write_bytes(data)
    return p


@pytest.fixture
def chart_pptx_path(tmp_path):
    pptx = pytest.importorskip("pptx", reason="python-pptx not installed")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C"]
    cd.add_series("S1", (1, 2, 3))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), cd
    )
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    return p


class TestDocTools:
    @pytest.mark.asyncio
    async def test_analyze_docx(self, docx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocAnalyzeTool().execute({"path": "doc.docx"}, ctx)
        assert not result.is_error, result.content
        info = json.loads(result.content)
        assert info["format"] == "docx"
        assert any("para" in item for item in info["outline"])

    @pytest.mark.asyncio
    async def test_apply_edits_xlsx_roundtrip(self, xlsx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocApplyEditsTool().execute(
            {
                "path": "book.xlsx",
                "edits": [{"action": "set_cell", "sheet": "Data", "cell": "B2", "value": 99}],
            },
            ctx,
        )
        assert not result.is_error, result.content
        summary = json.loads(result.content)
        assert summary["applied"] == 1
        assert summary["failed"] == 0

        # The edit is visible to a fresh analyze (in-place default output).
        check = await DocAnalyzeTool().execute({"path": summary["path"]}, ctx)
        assert "99" in check.content

    @pytest.mark.asyncio
    async def test_arrange_duplicate_slide(self, tmp_path):
        """DocArrange duplicates a slide (4 slides from 3), best-effort."""
        pytest.importorskip("pptx", reason="python-pptx not installed")
        ctx = ToolContext(working_dir=str(tmp_path))
        built = await DocBuildTool().execute(
            {
                "spec": {"slides": [{"title": "A"}, {"title": "B"}, {"title": "C"}]},
                "output": "deck.pptx",
            },
            ctx,
        )
        assert not built.is_error, built.content
        result = await DocArrangeTool().execute(
            {"path": "deck.pptx", "ops": [{"op": "duplicate", "target": 0, "to": 3}]},
            ctx,
        )
        assert not result.is_error, result.content
        summary = json.loads(result.content)
        assert summary["applied"] == 1 and summary["failed"] == 0
        info = await DocAnalyzeTool().execute({"path": summary["path"]}, ctx)
        assert len(json.loads(info.content)["slides"]) == 4

    @pytest.mark.asyncio
    async def test_edit_chart_retitle_and_data(self, chart_pptx_path, tmp_path):
        """Chart edits ride DocApplyEdits — a `chart` key routes them."""
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocApplyEditsTool().execute(
            {
                "path": "deck.pptx",
                "edits": [
                    {"chart": 0, "title": "Q3 Sales"},
                    {
                        "chart": 0,
                        "categories": ["Q1", "Q2", "Q3"],
                        "series": [{"name": "Rev", "values": [10, 20, 30]}],
                    },
                ],
            },
            ctx,
        )
        assert not result.is_error, result.content
        summary = json.loads(result.content)
        assert summary["applied"] == 2
        assert summary["failed"] == 0
        # The retitle + data change is visible to a fresh analyze.
        check = await DocAnalyzeTool().execute({"path": summary["path"]}, ctx)
        assert "Q3 Sales" in check.content
        assert "Rev" in check.content

    @pytest.mark.asyncio
    async def test_build_docx_from_markdown(self, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocBuildTool().execute(
            {"spec": "# Report\n\nBody **text**.\n\n- a\n- b", "output": "out.docx"},
            ctx,
        )
        assert not result.is_error, result.content
        summary = json.loads(result.content)
        assert (tmp_path / "out.docx").exists()
        # Round-trips through analyze.
        check = await DocAnalyzeTool().execute({"path": summary["path"]}, ctx)
        assert json.loads(check.content)["format"] == "docx"

    @pytest.mark.asyncio
    async def test_build_pptx_from_slide_spec(self, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocBuildTool().execute(
            {
                "spec": {"slides": [
                    {"layout": "title", "title": "Deck", "subtitle": "2026"},
                    {"layout": "content", "title": "Agenda", "bullets": ["A", "B"]},
                ]},
                "output": "deck.pptx",
            },
            ctx,
        )
        assert not result.is_error, result.content
        assert json.loads(result.content)["page_count"] == 2
        assert (tmp_path / "deck.pptx").exists()

    @pytest.mark.asyncio
    async def test_xml_read_lists_parts_and_reads_one(self, chart_pptx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        listing = await DocXmlReadTool().execute({"path": "deck.pptx"}, ctx)
        assert not listing.is_error, listing.content
        parts = json.loads(listing.content)["parts"]
        assert any("charts/chart1.xml" in p["part"] for p in parts)
        read = await DocXmlReadTool().execute(
            {"path": "deck.pptx", "part": "ppt/charts/chart1.xml"}, ctx
        )
        assert not read.is_error and "<c:ser>" in read.content

    @pytest.mark.asyncio
    async def test_xml_edit_recolors_chart_series(self, chart_pptx_path, tmp_path):
        """The real-world failure case: recolor bars — now a pure tool call."""
        pptx = pytest.importorskip("pptx")
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocXmlEditTool().execute(
            {
                "path": "deck.pptx",
                "part": "ppt/charts/chart1.xml",
                "edits": [{
                    "find": "</c:tx>",
                    "replace": (
                        "</c:tx><c:spPr><a:solidFill>"
                        '<a:srgbClr val="FF0000"/>'
                        "</a:solidFill></c:spPr>"
                    ),
                }],
            },
            ctx,
        )
        assert not result.is_error, result.content
        assert json.loads(result.content)["applied"] == 1
        prs = pptx.Presentation(str(chart_pptx_path))
        chart = next(s for sl in prs.slides for s in sl.shapes if s.has_chart).chart
        assert str(chart.series[0].format.fill.fore_color.rgb) == "FF0000"

    @pytest.mark.asyncio
    async def test_xml_edit_rejects_malformed_result(self, chart_pptx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocXmlEditTool().execute(
            {
                "path": "deck.pptx",
                "part": "ppt/charts/chart1.xml",
                "edits": [{"find": "</c:chartSpace>", "replace": "<broken"}],
            },
            ctx,
        )
        # Refused as engine feedback: nothing applied, doc still valid.
        assert not result.is_error
        summary = json.loads(result.content)
        assert summary["applied"] == 0 and summary["failed"] >= 1

    @pytest.mark.asyncio
    async def test_guide_root_map_with_executor_names(self, tmp_path):
        """DocGuide is the skill entry point: family map rendered with the
        EXECUTOR tool names (never the library verb names)."""
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocGuideTool().execute({}, ctx)
        assert not result.is_error, result.content
        assert "GENERATE" in result.content and "EDIT" in result.content
        assert "DocApplyEdits" in result.content
        assert "DocXmlEdit" in result.content
        assert "set_doc_text" not in result.content
        assert result.metadata["topics"]

    @pytest.mark.asyncio
    async def test_guide_topic_and_prefix(self, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        colors = await DocGuideTool().execute({"topic": "recipes.colors"}, ctx)
        assert "srgbClr" in colors.content and "DocXmlEdit" in colors.content
        recipes = await DocGuideTool().execute({"topic": "recipes"}, ctx)
        assert "COPY / MOVE / DELETE a slide" in recipes.content  # recipes.slides
        assert "RECOLOR" in recipes.content  # recipes.colors
        unknown = await DocGuideTool().execute({"topic": "zzz"}, ctx)
        assert not unknown.is_error and "GENERATE" in unknown.content

    def test_descriptions_stay_compact(self):
        """Progressive-disclosure contract: frontmatter tier stays small;
        the fat how-to lives behind DocGuide(topic)."""
        for name, cls in DOC_TOOL_CLASSES.items():
            desc = cls().description
            assert len(desc) <= 320, (
                f"{name} description grew to {len(desc)} chars — move "
                "detail into edit2docs agent_guide GUIDES instead"
            )

    def test_guide_registered_first(self):
        assert list(DOC_TOOL_CLASSES)[0] == "DocGuide"

    def test_llm_verbs_are_feature_gated(self):
        """Keyless hosts must never see DocGenerate/DocEdit — they advertise
        feature:docs_llm so progressive disclosure drops them."""
        assert DocGenerateTool().required_config_keys() == ["feature:docs_llm"]
        assert DocEditTool().required_config_keys() == ["feature:docs_llm"]

    @pytest.mark.asyncio
    async def test_xml_edit_creates_and_deletes_parts(self, chart_pptx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        # create a brand-new XML part (content type registered)
        created = await DocXmlEditTool().execute(
            {
                "path": "deck.pptx",
                "part": "ppt/slides/slide2.xml",
                "xml": (
                    await DocXmlReadTool().execute(
                        {"path": "deck.pptx", "part": "ppt/slides/slide1.xml"}, ctx
                    )
                ).content,
                "content_type": (
                    "application/vnd.openxmlformats-officedocument"
                    ".presentationml.slide+xml"
                ),
            },
            ctx,
        )
        assert not created.is_error, created.content
        listing = await DocXmlReadTool().execute({"path": "deck.pptx"}, ctx)
        names = [q["part"] for q in json.loads(listing.content)["parts"]]
        assert "ppt/slides/slide2.xml" in names
        # delete it again
        deleted = await DocXmlEditTool().execute(
            {"path": "deck.pptx", "part": "ppt/slides/slide2.xml", "delete": True},
            ctx,
        )
        assert not deleted.is_error, deleted.content
        listing = await DocXmlReadTool().execute({"path": "deck.pptx"}, ctx)
        names = [q["part"] for q in json.loads(listing.content)["parts"]]
        assert "ppt/slides/slide2.xml" not in names

    @pytest.mark.asyncio
    async def test_xml_edit_requires_exactly_one_mode(self, chart_pptx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocXmlEditTool().execute(
            {"path": "deck.pptx", "part": "ppt/charts/chart1.xml"}, ctx
        )
        assert result.is_error

    @pytest.mark.asyncio
    async def test_build_rejects_wrong_spec_type(self, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocBuildTool().execute(
            {"spec": "markdown-not-a-dict", "output": "x.pptx"}, ctx
        )
        assert result.is_error

    @pytest.mark.asyncio
    async def test_edit_chart_out_of_range_soft_fails(self, chart_pptx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocApplyEditsTool().execute(
            {"path": "deck.pptx", "edits": [{"chart": 9, "title": "nope"}]},
            ctx,
        )
        # Out-of-range is soft engine feedback, not a hard tool error.
        assert not result.is_error, result.content
        summary = json.loads(result.content)
        assert summary["applied"] == 0
        assert summary["failed"] == 1
        assert summary["results"][0]["status"] == "not_found"

    @pytest.mark.asyncio
    async def test_apply_edits_soft_fail_reports_status(self, xlsx_path, tmp_path):
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocApplyEditsTool().execute(
            {
                "path": "book.xlsx",
                "edits": [{"action": "set_cell", "sheet": "Nope", "cell": "A1", "value": 1}],
            },
            ctx,
        )
        assert not result.is_error  # soft-fail: statuses, not exceptions
        summary = json.loads(result.content)
        assert summary["applied"] == 0
        assert summary["failed"] == 1

    @pytest.mark.asyncio
    async def test_render_md_readable_content(self, docx_path, tmp_path):
        """DocPreview folded into DocRender: to='md' returns readable content."""
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocRenderTool().execute({"path": "doc.docx", "to": "md"}, ctx)
        assert not result.is_error, result.content
        payload = json.loads(result.content)
        assert payload["to"] == "md"
        md = Path(payload["paths"][0]).read_text(encoding="utf-8")
        assert "First paragraph" in md

    @pytest.mark.asyncio
    async def test_path_guard_blocks_escape(self, tmp_path):
        inner = tmp_path / "inner"
        inner.mkdir()
        ctx = ToolContext(working_dir=str(inner), allowed_paths=[str(inner)])
        result = await DocAnalyzeTool().execute({"path": "../../etc/passwd"}, ctx)
        assert result.is_error
        assert "Access denied" in result.content or "No such file" in result.content

    @pytest.mark.asyncio
    async def test_llm_verbs_require_key(self, docx_path, tmp_path, monkeypatch):
        monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
        ctx = ToolContext(working_dir=str(tmp_path))
        gen = await DocGenerateTool().execute(
            {"intent": "x", "output": "new.docx"}, ctx
        )
        assert gen.is_error
        assert "ANTHROPIC_API_KEY" in gen.content
        edit = await DocEditTool().execute(
            {"path": "doc.docx", "instruction": "x"}, ctx
        )
        assert edit.is_error
        assert "ANTHROPIC_API_KEY" in edit.content

    @pytest.mark.asyncio
    async def test_render_png_pages(self, docx_path, tmp_path):
        from xgen_agent_runtime.tools.built_in.doc_tools import DocRenderTool

        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocRenderTool().execute(
            {"path": "doc.docx", "to": "png", "out_dir": "prev"}, ctx
        )
        assert not result.is_error, result.content
        payload = json.loads(result.content)
        assert payload["page_count"] >= 1
        assert payload["paths"][0].endswith("page-1.png")
        assert (tmp_path / "prev" / "page-1.png").exists()

    @pytest.mark.asyncio
    async def test_render_pdf(self, docx_path, tmp_path):
        from xgen_agent_runtime.tools.built_in.doc_tools import DocRenderTool

        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocRenderTool().execute({"path": "doc.docx", "to": "pdf"}, ctx)
        assert not result.is_error, result.content
        payload = json.loads(result.content)
        assert payload["to"] == "pdf" and payload["paths"][0].endswith(".pdf")

    @pytest.mark.asyncio
    async def test_unsupported_extension_rejected(self, tmp_path):
        (tmp_path / "notes.txt").write_text("hi", encoding="utf-8")
        ctx = ToolContext(working_dir=str(tmp_path))
        result = await DocAnalyzeTool().execute({"path": "notes.txt"}, ctx)
        assert result.is_error
        assert "Unsupported document format" in result.content
